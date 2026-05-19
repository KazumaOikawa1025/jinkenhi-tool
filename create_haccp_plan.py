#!/usr/bin/env python3
"""五反田ふじ屋 HACCP衛生管理計画書 Excel生成スクリプト"""

import openpyxl
from openpyxl.styles import (
    Font, Alignment, Border, Side, PatternFill, numbers
)
from openpyxl.utils import get_column_letter

wb = openpyxl.Workbook()

# ── スタイル定義 ──
thin = Side(style="thin")
border_all = Border(top=thin, bottom=thin, left=thin, right=thin)
header_font = Font(name="游ゴシック", size=14, bold=True)
sub_header_font = Font(name="游ゴシック", size=11, bold=True)
normal_font = Font(name="游ゴシック", size=10)
small_font = Font(name="游ゴシック", size=9)
title_fill = PatternFill("solid", fgColor="1F4E79")
title_font = Font(name="游ゴシック", size=14, bold=True, color="FFFFFF")
section_fill = PatternFill("solid", fgColor="D6E4F0")
section_font = Font(name="游ゴシック", size=11, bold=True, color="1F4E79")
item_fill = PatternFill("solid", fgColor="F2F2F2")
center = Alignment(horizontal="center", vertical="center", wrap_text=True)
left_wrap = Alignment(horizontal="left", vertical="center", wrap_text=True)
left_top = Alignment(horizontal="left", vertical="top", wrap_text=True)


def apply_border(ws, row_start, row_end, col_start, col_end):
    for r in range(row_start, row_end + 1):
        for c in range(col_start, col_end + 1):
            ws.cell(r, c).border = border_all


def set_col_widths(ws, widths):
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


# ============================================================
# Sheet 1: 衛生管理計画書（一般衛生管理）
# ============================================================
ws1 = wb.active
ws1.title = "衛生管理計画（一般）"
ws1.sheet_properties.pageSetUpPr.fitToPage = True

set_col_widths(ws1, [5, 28, 40, 40, 40])

# タイトル
ws1.merge_cells("A1:E1")
c = ws1["A1"]
c.value = "衛生管理計画書（一般的衛生管理）"
c.font = title_font
c.fill = title_fill
c.alignment = center
ws1.row_dimensions[1].height = 36

# 店舗情報
ws1.merge_cells("A2:E2")
c = ws1["A2"]
c.value = "施設名称：五反田ふじ屋　　　　作成日：　　年　　月　　日　　　　作成者：　　　　　　　　"
c.font = sub_header_font
c.alignment = Alignment(horizontal="left", vertical="center")
ws1.row_dimensions[2].height = 28

# ヘッダ行
headers = ["No.", "管理項目", "いつ", "どのように", "問題があったとき"]
row = 3
for i, h in enumerate(headers, 1):
    cell = ws1.cell(row, i, h)
    cell.font = Font(name="游ゴシック", size=10, bold=True, color="FFFFFF")
    cell.fill = PatternFill("solid", fgColor="2E75B6")
    cell.alignment = center
ws1.row_dimensions[3].height = 26

# 一般衛生管理の項目
general_items = [
    (
        "1",
        "原材料の受入の確認",
        "原材料の搬入時",
        "・外観、におい、包装の状態、表示（期限、保存方法）を確認\n・問題がないものだけ受け入れる",
        "・返品し交換を求める\n・具体的な状況を記録する",
    ),
    (
        "2",
        "冷蔵・冷凍庫の温度の確認",
        "始業時・業務終了時",
        "・冷蔵庫：10℃以下\n・冷凍庫：-15℃以下\nを温度計で確認する",
        "・異常の原因を確認する\n・設定温度の再調整・故障の場合は修理を依頼\n・食材の状態によっては廃棄を判断",
    ),
    (
        "3",
        "交差汚染・二次汚染の防止",
        "作業中・洗浄時",
        "・器具（包丁・まな板等）は用途別に使い分ける\n・使用の都度、洗浄・消毒する\n・生肉・生魚を取り扱った後は手洗い",
        "・汚染された可能性のある食品は廃棄する\n・器具を洗浄・消毒する\n・手洗いを徹底する",
    ),
    (
        "4",
        "器具等の洗浄・消毒・殺菌",
        "使用後・業務終了時",
        "・器具・容器・ふきん等は使用後に洗浄・消毒する\n・業務終了後にまとめて洗浄・殺菌（85℃以上の熱湯又は次亜塩素酸ナトリウム等）",
        "・洗い直しをする\n・消毒・殺菌をやり直す\n・汚れた器具は使用しない",
    ),
    (
        "5",
        "トイレの洗浄・消毒",
        "始業前・営業中（定期的）",
        "・便座・ドアノブ・手洗い設備を洗浄・消毒\n・トイレ専用の履物を使用\n・消毒用アルコールの設置を確認",
        "・洗浄・消毒をやり直す\n・消耗品（石鹸・ペーパータオル等）を補充",
    ),
    (
        "6",
        "従業員の健康管理・\n衛生的作業着の着用等",
        "始業前",
        "・健康状態を確認（下痢・嘔吐・発熱等の有無）\n・手指の傷の有無を確認\n・清潔な作業着・帽子（三角巾）の着用",
        "・体調不良者は調理作業に従事させない\n・手指に傷がある場合は使い捨て手袋を着用\n・責任者に報告し対応を記録",
    ),
    (
        "7",
        "衛生的な手洗いの実施",
        "作業開始前・用便後\nメニュー変更時\n生肉等取扱い後",
        "・衛生的な手洗いを実施\n（水で流す→石鹸で洗う→水で流す→消毒）\n・使い捨てペーパータオルで拭く",
        "・手洗いをやり直す\n・手洗い方法を再指導する",
    ),
]

row = 4
for item in general_items:
    for i, val in enumerate(item, 1):
        cell = ws1.cell(row, i, val)
        cell.font = normal_font
        if i == 1:
            cell.alignment = center
        else:
            cell.alignment = left_wrap
    ws1.row_dimensions[row].height = 80
    row += 1

apply_border(ws1, 1, row - 1, 1, 5)

# ============================================================
# Sheet 2: 重要管理のポイント（メニュー別）
# ============================================================
ws2 = wb.create_sheet("重要管理（メニュー別）")
set_col_widths(ws2, [5, 22, 16, 40, 40, 40])

ws2.merge_cells("A1:F1")
c = ws2["A1"]
c.value = "衛生管理計画書（重要管理のポイント）"
c.font = title_font
c.fill = title_fill
c.alignment = center
ws2.row_dimensions[1].height = 36

ws2.merge_cells("A2:F2")
c = ws2["A2"]
c.value = "施設名称：五反田ふじ屋　　　　作成日：　　年　　月　　日　　　　作成者：　　　　　　　　"
c.font = sub_header_font
c.alignment = Alignment(horizontal="left", vertical="center")
ws2.row_dimensions[2].height = 28

# 説明
ws2.merge_cells("A3:F3")
c = ws2["A3"]
c.value = "※メニューを①加熱しないもの ②加熱するもの ③加熱後冷却するもの の3グループに分類し管理する"
c.font = Font(name="游ゴシック", size=9, color="FF0000")
c.alignment = left_wrap
ws2.row_dimensions[3].height = 22

headers2 = ["No.", "メニューグループ", "チェック方法", "管理のポイント", "問題があったとき", "対象メニュー例"]
row = 4
for i, h in enumerate(headers2, 1):
    cell = ws2.cell(row, i, h)
    cell.font = Font(name="游ゴシック", size=10, bold=True, color="FFFFFF")
    cell.fill = PatternFill("solid", fgColor="2E75B6")
    cell.alignment = center
ws2.row_dimensions[4].height = 26

menu_items = [
    (
        "①",
        "加熱しないもの\n（冷たいまま提供）",
        "目視・温度確認",
        "・冷蔵庫から出したらすみやかに提供\n・盛り付け時の衛生管理（手袋・清潔な器具の使用）\n・食材の鮮度・期限を確認",
        "・期限切れ・異常のある食材は廃棄\n・衛生的でない盛り付けはやり直し\n・具体的な状況を記録",
        "・刺身\n・サラダ\n・冷奴\n・漬物",
    ),
    (
        "②",
        "加熱するもの\n（加熱して提供）",
        "温度確認・目視",
        "・中心部まで十分に加熱する\n（中心温度75℃以上 1分以上）\n・温度計で確認 又は 火の通りを目視確認\n・加熱後はすみやかに提供",
        "・加熱不十分の場合は再加熱する\n・状況を記録する\n・食材に異常があれば廃棄",
        "・焼き鳥\n・揚げ物\n・煮物\n・焼き魚\n・味噌汁",
    ),
    (
        "③",
        "加熱後冷却し\n再加熱するもの\n又は冷却提供するもの",
        "温度確認・時間管理",
        "・加熱時：中心温度75℃以上 1分以上\n・速やかに冷却（30分以内に20℃付近、1時間以内に10℃以下）\n・再加熱時：中心温度75℃以上 1分以上\n・冷蔵保管中は10℃以下を確認",
        "・冷却が不十分な場合はやり直し\n・長時間常温放置した食品は廃棄\n・再加熱不十分の場合は再度加熱\n・状況を記録する",
        "・ポテトサラダ\n・煮込み料理の作り置き\n・タレ・ソース類\n・出汁の冷却保管",
    ),
]

row = 5
for item in menu_items:
    for i, val in enumerate(item, 1):
        cell = ws2.cell(row, i, val)
        cell.font = normal_font
        if i == 1:
            cell.alignment = center
        else:
            cell.alignment = left_wrap
    ws2.row_dimensions[row].height = 110
    row += 1

apply_border(ws2, 1, row - 1, 1, 6)

# ============================================================
# Sheet 3: 衛生管理の実施記録（日報）
# ============================================================
ws3 = wb.create_sheet("衛生管理記録（日報）")
set_col_widths(ws3, [5, 20, 6, 6, 6, 6, 6, 6, 6, 6, 6, 6, 6, 6, 6, 6, 6, 6, 6, 6, 6, 6, 6, 6, 6, 6, 6, 6, 6, 6, 6, 6, 6, 20])

ws3.merge_cells("A1:AH1")
c = ws3["A1"]
c.value = "衛生管理の実施記録　　　　施設名称：五反田ふじ屋　　　　　　年　　　月分"
c.font = title_font
c.fill = title_fill
c.alignment = center
ws3.row_dimensions[1].height = 36

# ヘッダ：管理項目 + 日付1〜31 + 特記事項
headers3 = ["No.", "管理項目"]
for d in range(1, 32):
    headers3.append(str(d))
headers3.append("特記事項")

row = 2
for i, h in enumerate(headers3, 1):
    cell = ws3.cell(row, i, h)
    cell.font = Font(name="游ゴシック", size=8, bold=True, color="FFFFFF")
    cell.fill = PatternFill("solid", fgColor="2E75B6")
    cell.alignment = center
ws3.row_dimensions[2].height = 22

record_items = [
    ("1", "原材料の受入の確認"),
    ("2", "冷蔵庫の温度（　　℃）"),
    ("3", "冷凍庫の温度（　　℃）"),
    ("4", "交差汚染・二次汚染の防止"),
    ("5", "器具等の洗浄・消毒・殺菌"),
    ("6", "トイレの洗浄・消毒"),
    ("7", "従業員の健康管理"),
    ("8", "衛生的な手洗いの実施"),
    ("9", "加熱温度の確認（　　℃）"),
    ("10", "冷却の確認"),
]

row = 3
for item in record_items:
    ws3.cell(row, 1, item[0]).font = small_font
    ws3.cell(row, 1).alignment = center
    ws3.cell(row, 2, item[1]).font = small_font
    ws3.cell(row, 2).alignment = left_wrap
    for d in range(3, 34):
        ws3.cell(row, d).alignment = center
        ws3.cell(row, d).font = small_font
    ws3.cell(row, 34).font = small_font
    ws3.cell(row, 34).alignment = left_wrap
    ws3.row_dimensions[row].height = 24
    row += 1

# 確認者・責任者欄
row += 1
ws3.merge_cells(f"A{row}:B{row}")
ws3.cell(row, 1, "確認者サイン").font = sub_header_font
ws3.cell(row, 1).alignment = center
ws3.row_dimensions[row].height = 30
for d in range(3, 34):
    ws3.cell(row, d).alignment = center

row += 1
ws3.merge_cells(f"A{row}:B{row}")
ws3.cell(row, 1, "責任者サイン").font = sub_header_font
ws3.cell(row, 1).alignment = center
ws3.row_dimensions[row].height = 30

apply_border(ws3, 2, row, 1, 34)

# 記入方法の注記
row += 2
ws3.merge_cells(f"A{row}:AH{row}")
c = ws3.cell(row, 1)
c.value = "【記入方法】  良好：○　　問題あり：×（特記事項欄に内容・対応を記録）　　該当なし：ー　　休業日：／"
c.font = Font(name="游ゴシック", size=9, color="FF0000")
c.alignment = left_wrap

# ============================================================
# Sheet 4: 緊急時対応計画
# ============================================================
ws4 = wb.create_sheet("緊急時対応計画")
set_col_widths(ws4, [5, 28, 50, 50])

ws4.merge_cells("A1:D1")
c = ws4["A1"]
c.value = "緊急時対応計画"
c.font = title_font
c.fill = title_fill
c.alignment = center
ws4.row_dimensions[1].height = 36

ws4.merge_cells("A2:D2")
c = ws4["A2"]
c.value = "施設名称：五反田ふじ屋　　　　作成日：　　年　　月　　日"
c.font = sub_header_font
c.alignment = Alignment(horizontal="left", vertical="center")
ws4.row_dimensions[2].height = 28

headers4 = ["No.", "想定される緊急事態", "対応手順", "連絡先"]
row = 3
for i, h in enumerate(headers4, 1):
    cell = ws4.cell(row, i, h)
    cell.font = Font(name="游ゴシック", size=10, bold=True, color="FFFFFF")
    cell.fill = PatternFill("solid", fgColor="2E75B6")
    cell.alignment = center
ws4.row_dimensions[3].height = 26

emergency_items = [
    (
        "1",
        "食中毒（疑い含む）が\n発生した場合",
        "①営業を直ちに中止する\n②保健所に速やかに届出・相談する\n③原因と思われる食品・原材料を保管する\n④従業員の健康状態を確認する\n⑤施設の消毒・洗浄を行う\n⑥保健所の指示に従う",
        "品川区保健所 生活衛生課\nTEL: 03-5742-9138\n\n最寄りの保健センター\nTEL:",
    ),
    (
        "2",
        "従業員の体調不良\n（嘔吐・下痢等）",
        "①該当者を直ちに調理作業から外す\n②嘔吐物等の処理は使い捨て手袋・マスク着用\n③次亜塩素酸ナトリウムで消毒\n④医療機関の受診を指示\n⑤症状が回復するまで調理作業に従事させない",
        "近隣医療機関:\nTEL:\n\n責任者連絡先:\nTEL:",
    ),
    (
        "3",
        "停電・設備故障\n（冷蔵庫等）",
        "①冷蔵庫・冷凍庫の開閉を最小限にする\n②食材の温度を確認する\n③復旧後、食材の状態を確認\n④温度逸脱した食材は廃棄を検討\n⑤修理業者に連絡する",
        "設備修理業者:\nTEL:\n\n電力会社:\nTEL:",
    ),
    (
        "4",
        "異物混入の\nクレームがあった場合",
        "①お客様に謝罪し状況を確認する\n②混入した異物を回収・保管する\n③原因を調査する\n④同一ロットの食品を確認する\n⑤再発防止策を検討・実施する\n⑥記録に残す",
        "責任者:\nTEL:\n\n必要に応じて保健所に相談",
    ),
]

row = 4
for item in emergency_items:
    for i, val in enumerate(item, 1):
        cell = ws4.cell(row, i, val)
        cell.font = normal_font
        if i == 1:
            cell.alignment = center
        else:
            cell.alignment = left_top
    ws4.row_dimensions[row].height = 120
    row += 1

apply_border(ws4, 1, row - 1, 1, 4)

# ============================================================
# Sheet 5: 施設情報・体制
# ============================================================
ws5 = wb.create_sheet("施設情報")
set_col_widths(ws5, [5, 22, 50])

ws5.merge_cells("A1:C1")
c = ws5["A1"]
c.value = "施設情報・衛生管理体制"
c.font = title_font
c.fill = title_fill
c.alignment = center
ws5.row_dimensions[1].height = 36

info_items = [
    ("1", "施設の名称", "五反田ふじ屋"),
    ("2", "施設の所在地", "東京都品川区西五反田　　丁目　　番　　号"),
    ("3", "営業者氏名", ""),
    ("4", "食品衛生責任者", ""),
    ("5", "営業許可番号", ""),
    ("6", "営業の種類", "飲食店営業"),
    ("7", "主な提供メニュー", ""),
    ("8", "営業時間", ""),
    ("9", "定休日", ""),
    ("10", "従業員数", "正社員：　　名　　アルバイト：　　名"),
    ("11", "衛生管理責任者", ""),
    ("12", "緊急連絡先", ""),
]

headers5 = ["No.", "項目", "内容"]
row = 2
for i, h in enumerate(headers5, 1):
    cell = ws5.cell(row, i, h)
    cell.font = Font(name="游ゴシック", size=10, bold=True, color="FFFFFF")
    cell.fill = PatternFill("solid", fgColor="2E75B6")
    cell.alignment = center
ws5.row_dimensions[2].height = 26

row = 3
for item in info_items:
    for i, val in enumerate(item, 1):
        cell = ws5.cell(row, i, val)
        cell.font = normal_font
        if i == 1:
            cell.alignment = center
        else:
            cell.alignment = left_wrap
    ws5.row_dimensions[row].height = 28
    row += 1

apply_border(ws5, 2, row - 1, 1, 3)

# ============================================================
# 印刷設定
# ============================================================
for ws in [ws1, ws2, ws4, ws5]:
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True

ws3.page_setup.orientation = "landscape"
ws3.page_setup.paperSize = ws3.PAPERSIZE_A3
ws3.page_setup.fitToWidth = 1
ws3.page_setup.fitToHeight = 1
ws3.sheet_properties.pageSetUpPr.fitToPage = True

# 保存
output_path = "/home/user/jinkenhi-tool/五反田ふじ屋_HACCP衛生管理計画書.xlsx"
wb.save(output_path)
print(f"作成完了: {output_path}")
